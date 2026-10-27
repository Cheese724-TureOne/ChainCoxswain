# -*- coding: utf-8 -*-
"""报告导出测试：xlsx / docx / pptx / html 可生成、可被办公库重新打开、内容可审计。

所有报告写进临时目录（monkeypatch report_gen.OUT_DIR），不污染产品 exports/。
"""
import os

import pytest

import report_gen
from conftest import PROJECT_ROOT  # noqa: F401


@pytest.fixture(scope="module")
def out_dir(tmp_path_factory):
    return str(tmp_path_factory.mktemp("exports"))


@pytest.fixture(autouse=True)
def _isolate_exports(out_dir, monkeypatch):
    monkeypatch.setattr(report_gen, "OUT_DIR", out_dir)


def test_xlsx_export_with_five_sheets(service, out_dir):
    res = service.export_report("xlsx")
    assert res["format"] == "xlsx"
    assert os.path.basename(res["path"]).startswith("ChainCoxswain_风险分析_")
    assert res["sizeBytes"] > 1000
    import openpyxl
    wb = openpyxl.load_workbook(res["path"])
    assert wb.sheetnames == ["库存风险清单", "概率与政策", "风险临界点", "干预方案", "数据来源与假设"]
    ws = wb["库存风险清单"]
    assert ws.max_row >= 8  # 标题行 + 表头 + 若干风险行
    ws4 = wb["数据来源与假设"]
    joined = " ".join(str(ws4.cell(row=i, column=1).value or "") for i in range(1, 20))
    assert "人工审核" in joined, "Excel 必须写明业务动作需人工审核"
    assert "演示数据" in joined


def test_docx_export_openable(service, out_dir):
    res = service.export_report("docx")
    assert os.path.basename(res["path"]).startswith("ChainCoxswain_分析报告_")
    from docx import Document
    doc = Document(res["path"])
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "供应链库存风险分析报告" in text
    assert "人工审核" in text
    assert len(doc.tables) >= 2, "Word 报告应含风险表与临界点表"


def test_pptx_export_openable(service, out_dir):
    res = service.export_report("pptx")
    assert os.path.basename(res["path"]).startswith("ChainCoxswain_汇报材料_")
    from pptx import Presentation
    prs = Presentation(res["path"])
    assert len(prs.slides._sldIdLst) >= 4
    texts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
    joined = "\n".join(texts)
    assert "供应链库存风险分析" in joined
    assert "人工审核" in joined
    assert "风险临界点" in joined


def test_html_export_contains_risk_table(service, out_dir):
    res = service.export_report("html")
    with open(res["path"], encoding="utf-8") as f:
        html = f.read()
    assert "<table>" in html
    assert "供应链库存风险分析报告" in html


def test_unknown_format_raises(service):
    with pytest.raises(ValueError):
        service.export_report("pdf")
